#!/usr/bin/env python3
"""Weekly JPX two-market margin balance (信用取引現在高) → Market Ledger rows.

v13.5.65 (stabilization item 5). The committed CSV
`ops/imports/jpx_two_market_credit_20020802_20260710.csv` ends on 2026-07-10;
later weeks were meant to arrive through `/api/argus/admin/market-ledger/import`
but nothing produced them, so the conditioning engine's 45-day window dropped
the credit features (信用倍率 / 売り残高) from every forecast after August 24.

This tool fetches the official weekly workbook for each Friday after `--since`
(`https://www.jpx.co.jp/markets/statistics-equities/margin/tvdivq0000001rk9-att/
mtseisanYYYYMMDD00.xls`), reads the two-market TOTAL value columns (委託+自己,
金額, 百万円 → JPY) exactly as the committed CSV does, and emits rows in the
ledger's CSV contract. With `--import` it posts them to the admin import route
(dry run, then commit) and reads the ledger back. Weeks whose workbook is not
published (404) are reported as gaps, never filled.

New imports use the actual successful fetch time as availableFrom. The
publication time is not guessed from the week end. Stored observations retain
their original availability and are never rewritten by this collector.
The format from 2026-09-25 is XLSX; older weeks retain the XLS adapter.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import math
import zipfile
import json
import os
import re
import sys
import urllib.error
import urllib.request
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser
from urllib.parse import urljoin, urlparse
from typing import Any, Dict, Iterable, List, Optional, Sequence

URL_TEMPLATE = ("https://www.jpx.co.jp/markets/statistics-equities/margin/"
                "tvdivq0000001rk9-att/mtseisan{ymd}00.xls")
NEW_FORMAT_FROM = "2026-09-25"
NEW_URL_TEMPLATE = ("https://www.jpx.co.jp/markets/statistics-equities/margin/"
                    "tvdivq0000001rk9-att/{ymd}_mtcurrent.xlsx")
LISTING_URL = "https://www.jpx.co.jp/markets/statistics-equities/margin/04.html"
MAX_WORKBOOK_BYTES = 4 * 1024 * 1024
MAX_LISTING_BYTES = 512 * 1024
CSV_COLUMNS = ("seriesId", "periodEnd", "publishedAt", "availableFrom",
               "observedAt", "value", "unit", "source", "sourceKind", "status", "metadata")
SERIES = {"short": "credit.short_balance", "long": "credit.long_balance"}
MILLION = 1_000_000
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TITLE_RE = re.compile(r"(\d{4})/(\d{1,2})/(\d{1,2})")


def fridays_after(since: str, today: Optional[date] = None) -> List[str]:
    start = date.fromisoformat(since)
    end = today or datetime.now(timezone(timedelta(hours=9))).date()
    if start > end or (end - start).days > 2 * 366:
        raise ValueError("jpx_period_range_invalid")
    out: List[str] = []
    day = start + timedelta(days=1)
    while day <= end:
        if day.weekday() == 4:
            # A closed Friday uses the week's last actual TSE session.
            import argus_market_clock as clock
            period = day
            try:
                while period.weekday() > 0 and not clock.canonical_trading_day(clock.JP_EQUITY, period):
                    period -= timedelta(days=1)
                if not clock.canonical_trading_day(clock.JP_EQUITY, period):
                    day += timedelta(days=1)
                    continue
            except clock.CalendarUnavailableError:
                period = day  # Never invent a holiday outside known coverage.
            if period > start:
                out.append(period.isoformat())
        day += timedelta(days=1)
    return out


def parse_sheet(grid: Sequence[Sequence[Any]]) -> Dict[str, Any]:
    """Pure: the two-market TOTAL value row of the JPX weekly layout.

    grid[r][c] are cell values. The title row carries the period date; the
    header row containing 「合計」 fixes the column of the total block whose
    first two value columns are 売残高 (sales) and 買残高 (purchases) with a
    weekly-change column between them; the 二市場計 block has a 株数 row and
    a 金額 row (百万円)."""
    text = [[str(cell if cell is not None else "") for cell in row] for row in grid]
    period = None
    for row in text[:3]:
        for cell in row:
            match = TITLE_RE.search(cell)
            if match:
                period = date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()
                break
        if period:
            break
    if not period:
        raise ValueError("jpx_period_not_found")
    if any("Total Outstanding Margin Trading" in cell for row in text for cell in row):
        return parse_current_sheet(grid, text, period)
    total_col = None
    for row in text:
        for index, cell in enumerate(row):
            if "合" in cell and "計" in cell and "Total" in cell:
                total_col = index
                break
        if total_col is not None:
            break
    if total_col is None:
        raise ValueError("jpx_total_column_not_found")
    value_row = None
    for index, row in enumerate(text):
        if any("二市場計" in cell for cell in row):
            for candidate in range(index, min(index + 3, len(text))):
                if any("金額" in cell for cell in text[candidate]):
                    value_row = candidate
                    break
            break
    if value_row is None:
        raise ValueError("jpx_value_row_not_found")

    def number(cell: Any) -> float:
        value = float(str(cell).replace(",", "").replace("▲", "-").strip())
        if not math.isfinite(value) or value <= 0:
            raise ValueError("jpx_non_positive_value")
        return value

    short_million = number(grid[value_row][total_col])
    long_million = number(grid[value_row][total_col + 2])
    return {"periodEnd": period,
            "shortJpy": int(round(short_million * MILLION)),
            "longJpy": int(round(long_million * MILLION))}


def parse_current_sheet(grid, text, period):
    """Read two-market amounts, never the separate Tokyo or share columns."""
    market_cols = [i for row in text[:8] for i, cell in enumerate(row)
                   if "Tokyo&Nagoya" in cell and "二市場" in cell]
    if len(market_cols) != 1:
        raise ValueError("jpx_two_market_column_not_found")
    col = market_cols[0]
    headers = text[6] if len(text) > 6 else []
    if len(headers) <= col + 2 or "Sales" not in headers[col] or "Purchases" not in headers[col + 2]:
        raise ValueError("jpx_amount_columns_not_found")
    total_rows = [i for i, row in enumerate(text)
                  if any("Total Outstanding Margin Trading" in cell for cell in row)]
    if len(total_rows) != 1:
        raise ValueError("jpx_total_row_not_found")
    row = total_rows[0] + 1
    if row >= len(text) or not any("金額" in cell and "Val." in cell for cell in text[row]):
        raise ValueError("jpx_total_amount_row_not_found")
    if not any("百万円" in cell for cells in text[:6] for cell in cells):
        raise ValueError("jpx_amount_unit_not_found")
    values = []
    for c in (col, col + 2):
        value = float(grid[row][c])
        if not math.isfinite(value) or value <= 0:
            raise ValueError("jpx_non_positive_value")
        values.append(int(round(value * MILLION)))
    return {"periodEnd": period, "shortJpy": values[0], "longJpy": values[1]}


def load_workbook_grid(payload: bytes) -> List[List[Any]]:
    if len(payload) > MAX_WORKBOOK_BYTES:
        raise ValueError("jpx_workbook_too_large")
    if payload.startswith(b"PK"):
        # The spreadsheet library is loaded only by this scheduled tool.
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            if len(archive.infolist()) > 100 or sum(x.file_size for x in archive.infolist()) > 16 * 1024 * 1024:
                raise ValueError("jpx_workbook_too_large")
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            if sheet.max_row > 200 or sheet.max_column > 64:
                raise ValueError("jpx_workbook_shape_changed")
            return [[cell.strftime("%Y/%m/%d") if isinstance(cell, (date, datetime)) else cell
                     for cell in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            book.close()
    import xlrd
    if payload.lstrip().lower().startswith((b"<!doc", b"<html")):
        raise ValueError("jpx_not_a_workbook")
    book = xlrd.open_workbook(file_contents=payload)
    sheet = book.sheet_by_index(0)
    if sheet.nrows > 200 or sheet.ncols > 64:
        raise ValueError("jpx_workbook_shape_changed")
    return [[sheet.cell_value(r, c) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)]


def build_rows(parsed: Dict[str, Any], *, url: str, sha256: str,
               observed_at: str) -> List[Dict[str, Any]]:
    observed = datetime.fromisoformat(observed_at.replace("Z", "+00:00"))
    if observed.tzinfo is None or observed.date() < date.fromisoformat(parsed["periodEnd"]):
        raise ValueError("jpx_invalid_fetch_time")
    stamp = observed_at
    source = (f"JPX official | {url} | sha256={sha256} | "
              "publication=weekly_final | availability=successful_fetch")
    return [{
        "seriesId": SERIES[key], "periodEnd": parsed["periodEnd"],
        "publishedAt": "", "availableFrom": stamp, "observedAt": observed_at,
        "value": parsed["shortJpy"] if key == "short" else parsed["longJpy"],
        "unit": "JPY", "source": source, "sourceKind": "official", "status": "live",
    } for key in ("short", "long")]


def rows_to_csv(rows: Iterable[Dict[str, Any]]) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(CSV_COLUMNS), lineterminator="\n")
    writer.writeheader()
    for row in rows:
        values = {key: row.get(key, "") for key in CSV_COLUMNS}
        if isinstance(values.get("metadata"), dict):
            values["metadata"] = json.dumps(values["metadata"], ensure_ascii=False, allow_nan=False, separators=(",", ":"))
        writer.writerow(values)
    return buffer.getvalue()


def fetch(url: str, timeout: int = 60) -> Optional[bytes]:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (ARGUS jpx-credit-weekly)"})
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = response.read(MAX_WORKBOOK_BYTES + 1)
            if len(payload) > MAX_WORKBOOK_BYTES:
                raise ValueError("jpx_workbook_too_large")
            return payload
    except urllib.error.HTTPError as error:
        if error.code == 404:
            return None
        raise



def publication_due(period: str, now: datetime) -> Optional[bool]:
    """Second actual TSE session after the held week, 16:00 JST.

    The nominal publication schedule is not a historical receipt timestamp.
    Outside the authoritative calendar coverage the due state is unknown.
    """
    from argus_credit_publication import publication_at
    due = publication_at(period)
    return None if due is None else now >= due


def workbook_links(html: str) -> Dict[str, str]:
    """Only official two-market workbooks; never PDF, general-margin or offsite."""
    class Links(HTMLParser):
        def __init__(self):
            super().__init__(); self.href = None; self.text = []; self.links = {}
        def handle_starttag(self, tag, attrs):
            if tag == 'a':
                self.href = dict(attrs).get('href'); self.text = []
        def handle_data(self, data):
            if self.href is not None:
                self.text.append(data)
        def handle_endtag(self, tag):
            if tag != 'a' or self.href is None:
                return
            href, title = self.href, ''.join(self.text)
            self.href = None
            url = urljoin(LISTING_URL, href)
            parsed = urlparse(url)
            if (parsed.scheme != 'https' or parsed.netloc != 'www.jpx.co.jp'
                    or parsed.query or parsed.fragment
                    or not parsed.path.startswith('/markets/statistics-equities/margin/tvdivq0000001rk9-att/')
                    or not parsed.path.endswith(('.xls', '.xlsx'))
                    or 'mtgaisan' in parsed.path):
                return
            matched = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', title)
            if not matched:
                matched = re.search(r'(\d{4})(\d{2})(\d{2})', parsed.path.rsplit('/', 1)[-1])
            if not matched:
                raise ValueError('jpx_listing_period_missing')
            period = date(*map(int, matched.groups())).isoformat()
            if period in self.links and self.links[period] != url:
                raise ValueError('jpx_listing_period_ambiguous')
            self.links[period] = url
    parser = Links(); parser.feed(html); parser.close()
    if not parser.links:
        raise ValueError('jpx_listing_workbooks_missing')
    return parser.links


def fetch_listing() -> Dict[str, str]:
    request = urllib.request.Request(LISTING_URL, headers={'User-Agent': 'Mozilla/5.0 (ARGUS jpx-credit-weekly)'})
    with urllib.request.urlopen(request, timeout=30) as response:
        body = response.read(MAX_LISTING_BYTES + 1)
        if len(body) > MAX_LISTING_BYTES:
            raise ValueError('jpx_listing_too_large')
    return workbook_links(body.decode('utf-8'))

def collect(since: str, *, today: Optional[date] = None,
            fetcher=fetch, now_iso: Optional[str] = None, listing_fetcher=fetch_listing,
            valuation_since: Optional[str] = None) -> Dict[str, Any]:
    observed = now_iso or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    clock_now = datetime.fromisoformat(observed.replace('Z', '+00:00'))
    rows: List[Dict[str, Any]] = []
    gaps: List[str] = []
    failures: List[str] = []
    fetched: List[str] = []
    starts = [since] + ([valuation_since] if valuation_since is not None else [])
    for start in starts:
        fridays_after(start, today)  # Validate each independent cursor, not only the minimum.
    links = listing_fetcher() if listing_fetcher is not None else {}
    for friday in fridays_after(min(starts), today):
        template = NEW_URL_TEMPLATE if friday >= NEW_FORMAT_FROM else URL_TEMPLATE
        url = links.get(friday) or template.format(ymd=friday.replace("-", ""))
        due = friday in links or publication_due(friday, clock_now) is True
        payload = fetcher(url)
        if payload is None:
            gaps.append(friday)
            if due:
                failures.append(friday + ":published_workbook_unavailable")
            continue
        try:
            grid = load_workbook_grid(payload)
            parsed = parse_sheet(grid)
        except ValueError as error:
            gaps.append(f"{friday}:{error}")
            failures.append(friday + ":workbook_unreadable")
            continue
        if parsed["periodEnd"] != friday:
            gaps.append(f"{friday}:period_mismatch:{parsed['periodEnd']}")
            failures.append(friday + ":period_mismatch")
            continue
        observed = now_iso or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        digest = hashlib.sha256(payload).hexdigest()
        if friday > since:
            rows.extend(build_rows(parsed, url=url, sha256=digest, observed_at=observed))
        if valuation_since is not None and friday > valuation_since:
            from argus_jpx_credit_valuation import extract, ledger_observation
            try:
                calculation = extract(grid)
                rows.append(ledger_observation(calculation, url=url, sha256=digest,
                                               received_at=observed))
            except ValueError:
                failures.append(friday + ":valuation_inputs_unreadable")
        fetched.append(friday)
    return {"since": since, "fetched": fetched, "gaps": gaps, "failures": failures, "rows": rows,
            "csv": rows_to_csv(rows), "observedAt": observed}


def post_json(url: str, body: Dict[str, Any], token: str, timeout: int = 600) -> Dict[str, Any]:
    data = json.dumps(body).encode("utf-8")
    request = urllib.request.Request(url, data=data, method="POST", headers={
        "Content-Type": "application/json", "X-ARGUS-ADMIN-TOKEN": token})
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8"))


def import_rows(csv_text: str, *, backend: str, token: str,
                expected_newest: Optional[str] = None) -> Dict[str, Any]:
    """Dry run, commit, read back. The commit rebuilds the ledger and can
    outlive the proxy's response window (2026-09-08: the rows landed but the
    client saw a transport error) — so a transport error on the commit is
    settled by the read-back: the import is a success when the ledger now
    holds the newest period we sent, and a failure otherwise."""
    endpoint = backend.rstrip("/") + "/api/argus/admin/market-ledger/import"
    dry = post_json(endpoint, {"csv": csv_text, "dryRun": True}, token)
    if not dry.get("ok") or dry.get("errors"):
        return {"ok": False, "stage": "dry_run", "errors": (dry.get("errors") or [])[:5]}
    try:
        commit = post_json(endpoint, {"csv": csv_text, "dryRun": False}, token)
    except Exception as error:                     # transport only; verified below
        commit = {"ok": None, "transportError": type(error).__name__}
    if commit.get("ok") is False or commit.get("errors"):
        return {"ok": False, "stage": "commit", "errors": (commit.get("errors") or [])[:5]}
    readback = ledger_newest_credit(backend, token=token)
    held = min((str(readback.get(sid, {}).get("periodEnd") or "") for sid in SERIES.values()), default="")
    if expected_newest and held < expected_newest:
        return {"ok": False, "stage": "readback", "expectedNewest": expected_newest,
                "ledger": readback, "transportError": commit.get("transportError")}
    import argus_market_ledger as market_ledger
    from argus_jpx_credit_valuation import ledger_observation_digest
    try:
        expected = market_ledger.parse_csv(csv_text)
        required = set(SERIES.values()) if expected_newest else set()
        required |= {row["seriesId"] for row in expected}
        if not required.issubset(readback):
            raise ValueError("jpx_readback_series_missing")
        for row in expected:
            sid = row["seriesId"]
            if sid not in (*SERIES.values(), "credit.valuation_loss_pct"):
                raise ValueError("jpx_import_series_invalid")
            row["value"] = None if row["value"] == "" else float(row["value"])
            seen = readback.get(sid, {})
            observations = [{"periodEnd": seen.get("periodEnd"), "value": seen.get("latestValue"),
                             "availableFrom": seen.get("availableFrom")}, *seen.get("history", [])]
            digest = ledger_observation_digest(row) if sid == "credit.valuation_loss_pct" else None
            matches = [x for x in observations if x.get("periodEnd") == row["periodEnd"]
                       and x.get("value") == row["value"]
                       and x.get("availableFrom") == row["availableFrom"]]
            if digest:
                matches = [x for x in matches if x.get("unit") == "percent"
                           and x.get("calculationDigest") == digest]
                for item in matches:
                    if item.get("auditedObservation") is not None:
                        if ledger_observation_digest(item["auditedObservation"]) != digest:
                            raise ValueError("jpx_readback_calculation_mismatch")
            if not matches:
                raise ValueError("jpx_readback_content_mismatch")
    except (KeyError, TypeError, ValueError):
        return {"ok": False, "stage": "readback_content", "ledger": readback}
    return {"ok": True, "stage": "verified_import", "importId": commit.get("importId"),
            "rowCount": len(commit.get("preview") or []),
            "settledByReadback": commit.get("ok") is None,
            "transportError": commit.get("transportError"), "ledger": readback}


def ledger_newest_credit(backend: str, *, token: str) -> Dict[str, Any]:
    if not token:
        raise ValueError("jpx_admin_token_missing")
    request = urllib.request.Request(backend.rstrip("/") + "/api/argus/market-ledger",
                                    headers={"X-ARGUS-ADMIN-TOKEN": token})
    with urllib.request.urlopen(request, timeout=180) as response:
        doc = json.loads(response.read().decode("utf-8"))
    out: Dict[str, Any] = {}
    for row in doc.get("table") or []:
        if row.get("seriesId") in (*SERIES.values(), "credit.valuation_loss_pct"):
            out[row["seriesId"]] = {"periodEnd": row.get("periodEnd"),
                                    "latestValue": row.get("latestValue"),
                                    "availableFrom": row.get("availableFrom"),
                                    "history": row.get("history") or [],
                                    "acquisition": row.get("acquisition"), "sourceKind": row.get("sourceKind")}
    return out


def valuation_cursor(table, fallback="2026-07-10"):
    from argus_jpx_credit_valuation import audited_ledger_observation, ledger_observation_digest
    for row in table:
        if row.get("seriesId") != "credit.valuation_loss_pct":
            continue
        valid = []
        for item in row.get("history") or []:
            audit = item.get("auditedObservation")
            if (audit and audited_ledger_observation(audit)
                    and item.get("calculationDigest") == ledger_observation_digest(audit)
                    and item.get("periodEnd") == audit.get("periodEnd")):
                valid.append(audit["periodEnd"])
        if valid:
            return max(fallback, max(valid))
    return fallback


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--since", default="2026-07-10", help="last period already held (exclusive)")
    parser.add_argument("--out", default=None, help="write the CSV rows here")
    parser.add_argument("--import", dest="do_import", action="store_true")
    parser.add_argument("--backend", default="https://argus-backend-3j2m.onrender.com")
    parser.add_argument("--token-env", default="ARGUS_ADMIN_TOKEN")
    parser.add_argument("--summary", default=None, help="write a JSON summary here")
    parser.add_argument("--valuation-since", default=None, help="last audited valuation period held (exclusive); independent from credit balances")
    parser.add_argument("--valuation-monthly-backfill", action="store_true", help="one-time missing audited periods from official monthly reports; balances are not imported")
    args = parser.parse_args(argv)
    try:
        if args.valuation_monthly_backfill:
            from scripts.jpx_credit_valuation_monthly import collect as monthly_collect
            held = ledger_newest_credit(args.backend, token=os.environ.get(args.token_env, ""))
            rows = monthly_collect([dict(row, seriesId=sid) for sid, row in held.items()])
            result = {"fetched": [r["periodEnd"] for r in rows], "gaps": [], "failures": [],
                      "rows": rows, "csv": rows_to_csv(rows)}
        else:
            options = {"valuation_since": args.valuation_since} if args.valuation_since is not None else {}
            result = collect(args.since, **options)
    except Exception as error:
        # Do not expose raw responses, URLs or authenticated records in public logs.
        print(json.dumps({'ok': False, 'stage': 'acquisition', 'errorClass': type(error).__name__}))
        return 1
    summary: Dict[str, Any] = {"since": args.since, "fetched": result["fetched"],
                               "gaps": result["gaps"], "failures": result.get("failures", []), "rowCount": len(result["rows"]),
                               "newestPeriod": (result["fetched"] or [None])[-1]}
    if args.out:
        with open(args.out, "w", encoding="utf-8") as handle:
            handle.write(result["csv"])
    if args.do_import and result["rows"]:
        token = os.environ.get(args.token_env, "")
        if not token:
            summary["import"] = {"ok": False, "stage": "token_missing"}
        else:
            try:
                summary["import"] = import_rows(result["csv"], backend=args.backend, token=token,
                    expected_newest=max((r["periodEnd"] for r in result["rows"]
                                         if r["seriesId"] in SERIES.values()), default=None))
                summary["ledger"] = (summary["import"] or {}).get("ledger") or ledger_newest_credit(args.backend, token=token)
            except Exception as error:
                summary["import"] = {"ok": False, "stage": "authenticated_import",
                                     "errorClass": type(error).__name__}
    if result["gaps"]:
        # A fixed warning does not expose rows or protected readback.
        print("::warning title=jpx-credit-weekly::取得待ちまたは書式を確認できない週があります")
    if args.summary:
        with open(args.summary, "w", encoding="utf-8") as handle:
            json.dump(summary, handle, ensure_ascii=False, indent=1)
    if result.get("failures"):
        print(json.dumps({"ok": False, "stage": "published_weeks_incomplete"}))
        return 1
    print(json.dumps({"ok": not args.do_import or not result["rows"] or bool((summary.get("import") or {}).get("ok")),
                      "stage": (summary.get("import") or {}).get("stage", "no_new_rows" if not result["rows"] else "complete")}, ensure_ascii=False))
    if args.do_import and result["rows"] and not (summary.get("import") or {}).get("ok"):
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
