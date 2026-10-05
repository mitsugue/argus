"""v13.5.65 — weekly JPX two-market credit rows for the Market Ledger."""
import csv
import io
from datetime import date

import scripts.jpx_credit_weekly as jw

GRID = [
    ["信用取引現在高（2026/8/28現在）"] + [""] * 14,
    [""] * 15, [""] * 15, [""] * 15,
    ["", "", "", "委　　託 Customer", "", "", "", "自　　己 Proprietary", "", "", "", "合　　計 Total", "", "", ""],
    ["", "", "", "売残高\nSales", "前週比", "買残高\nPurchases", "前週比", "売残高", "前週比", "買残高", "前週比", "売残高\nSales", "前週比", "買残高\nPurchases", "前週比"],
    ["", "二市場計\nTotal", "株数Shs.", 352942.0, -40978.0, 3478719.0, -25502.0, 108556.0, -41.0, 1267.0, -82.0, 461498.0, -41019.0, 3479986.0, -25584.0],
    ["", "", "金額Val.", 669333.0, -75990.0, 6526845.0, 46607.0, 155713.0, 4855.0, 2518.0, 213.0, 825046.0, -71135.0, 6529363.0, 46820.0],
]


def test_parse_reads_the_two_market_total_value_columns_in_yen():
    parsed = jw.parse_sheet(GRID)
    assert parsed == {"periodEnd": "2026-08-28", "shortJpy": 825_046_000_000,
                      "longJpy": 6_529_363_000_000}


def test_rows_use_actual_fetch_availability_without_guessing_publication():
    rows = jw.build_rows(jw.parse_sheet(GRID), url="https://example.test/x.xls",
                         sha256="ab" * 32, observed_at="2026-09-08T00:00:00Z")
    assert [r["seriesId"] for r in rows] == ["credit.short_balance", "credit.long_balance"]
    assert rows[0]["periodEnd"] == "2026-08-28"
    assert rows[0]["availableFrom"] == rows[0]["observedAt"] == "2026-09-08T00:00:00Z"
    assert rows[0]["publishedAt"] == ""
    assert rows[0]["value"] == 825_046_000_000 and rows[1]["value"] == 6_529_363_000_000
    assert rows[0]["unit"] == "JPY" and rows[0]["sourceKind"] == "official" and rows[0]["status"] == "live"
    assert "sha256=" + "ab" * 32 in rows[0]["source"] and "publication=weekly_final" in rows[0]["source"]
    text = jw.rows_to_csv(rows)
    parsed = list(csv.DictReader(io.StringIO(text)))
    assert list(parsed[0].keys()) == list(jw.CSV_COLUMNS)


def test_fridays_after_and_gaps_are_reported_not_filled():
    assert jw.fridays_after("2026-07-10", today=date(2026, 7, 31)) == ["2026-07-17", "2026-07-24", "2026-07-31"]
    served = {"2026-07-31"}

    def fetcher(url):
        for day in served:
            if day.replace("-", "") in url:
                return b"workbook"
        return None
    grid_by_url = {}

    def loader(payload):
        return GRID_0731
    GRID_0731 = [row[:] for row in GRID]
    GRID_0731[0][0] = "信用取引現在高（2026/7/31現在）"
    original = jw.load_workbook_grid
    jw.load_workbook_grid = loader
    try:
        result = jw.collect("2026-07-10", today=date(2026, 7, 31), fetcher=fetcher, listing_fetcher=None,
                            now_iso="2026-09-08T00:00:00Z")
    finally:
        jw.load_workbook_grid = original
    assert result["fetched"] == ["2026-07-31"]
    assert result["gaps"] == ["2026-07-17", "2026-07-24"]
    assert len(result["rows"]) == 2 and result["rows"][0]["periodEnd"] == "2026-07-31"


def test_parse_refuses_a_workbook_without_the_total_block():
    broken = [row[:] for row in GRID]
    broken[4] = [""] * 15
    try:
        jw.parse_sheet(broken)
    except ValueError as error:
        assert str(error) == "jpx_total_column_not_found"
    else:
        raise AssertionError("expected jpx_total_column_not_found")


def test_a_transport_error_on_commit_is_settled_by_the_ledger_read_back(monkeypatch):
    calls = []

    def fake_post(url, body, token, timeout=600):
        calls.append(body["dryRun"])
        if body["dryRun"]:
            return {"ok": True, "errors": [], "preview": [1, 2]}
        raise TimeoutError("proxy closed the connection")
    monkeypatch.setattr(jw, "post_json", fake_post)
    monkeypatch.setattr(jw, "ledger_newest_credit",
                        lambda backend, *, token: confirmed_readback())
    result = jw.import_rows(credit_csv(), backend="https://x", token="t", expected_newest="2026-08-28")
    assert result["ok"] is True and result["settledByReadback"] is True
    assert "TimeoutError" in result["transportError"]
    # …but not when the ledger does not hold what we sent
    monkeypatch.setattr(jw, "ledger_newest_credit",
                        lambda backend, *, token: {"credit.short_balance": {"periodEnd": "2026-07-10"},
                                         "credit.long_balance": {"periodEnd": "2026-07-10"}})
    result = jw.import_rows(credit_csv(), backend="https://x", token="t", expected_newest="2026-08-28")
    assert result["ok"] is False and result["stage"] == "readback"


def current_grid():
    rows = [[""] * 12 for _ in range(20)]
    rows[0][1] = "2026/9/25"
    rows[4][11] = "(単位：千株、百万円）"
    rows[5][4] = "二市場 Tokyo&Nagoya"
    rows[5][8] = "東京 Tokyo"
    rows[6][4], rows[6][6] = "売残高 Sales", "買残高 Purchases"
    rows[13][1] = "信用取引残高合計 Total Outstanding Margin Trading"
    rows[13][3] = "株数Shs."
    rows[13][4], rows[13][6] = 111, 222
    rows[14][3] = "金額Val."
    rows[14][4], rows[14][6] = 1000, 7000
    rows[14][8], rows[14][10] = 999, 6999
    return rows


def test_new_layout_uses_two_market_amounts_and_never_tokyo_or_shares():
    import pytest
    grid = current_grid()
    assert jw.parse_sheet(grid) == {"periodEnd": "2026-09-25", "shortJpy": 1000000000, "longJpy": 7000000000}
    for where, value in (((5, 4), "東京 Tokyo"), ((6, 6), "Weekly change"), ((14, 3), "株数Shs."), ((4, 11), "千株"), ((14, 4), float("nan"))):
        bad = [row[:] for row in grid]
        bad[where[0]][where[1]] = value
        with pytest.raises(ValueError):
            jw.parse_sheet(bad)


def test_xlsx_reader_preserves_excel_dates_and_closes_workbook():
    import openpyxl
    import datetime
    book = openpyxl.Workbook()
    for row in current_grid(): book.active.append(row)
    book.active["B1"] = datetime.datetime(2026, 9, 25)
    buffer = io.BytesIO(); book.save(buffer); book.close()
    parsed = jw.parse_sheet(jw.load_workbook_grid(buffer.getvalue()))
    assert parsed["periodEnd"] == "2026-09-25"
    assert parsed["shortJpy"] == 1000000000


def test_readback_carries_admin_token_and_rejects_missing_token(monkeypatch):
    import json
    import pytest
    captured = []
    def open_read(request, timeout):
        captured.append((request, timeout))
        return io.BytesIO(json.dumps({"table": [{"seriesId": "credit.short_balance", "periodEnd": "2026-09-25", "latestValue": 1000}]}).encode())
    monkeypatch.setattr(jw.urllib.request, "urlopen", open_read)
    out = jw.ledger_newest_credit("https://backend.test", token="fixture-token")
    assert out["credit.short_balance"]["periodEnd"] == "2026-09-25"
    assert captured[0][0].get_header("X-argus-admin-token") == "fixture-token"
    with pytest.raises(ValueError, match="jpx_admin_token_missing"):
        jw.ledger_newest_credit("https://backend.test", token="")
    assert len(captured) == 1


def test_missing_buy_series_cannot_prove_a_completed_import(monkeypatch):
    monkeypatch.setattr(jw, "post_json", lambda *a, **k: {"ok": True})
    monkeypatch.setattr(jw, "ledger_newest_credit", lambda backend, *, token: {"credit.short_balance": {"periodEnd": "2026-09-25"}})
    result = jw.import_rows("csv", backend="https://backend.test", token="t", expected_newest="2026-09-25")
    assert result["ok"] is False and result["stage"] == "readback"


def test_current_weeks_request_the_official_xlsx_name_and_old_weeks_keep_xls(monkeypatch):
    calls = []
    def fetcher(url):
        calls.append(url)
        return None
    out = jw.collect("2026-09-11", today=date(2026, 9, 25), fetcher=fetcher, listing_fetcher=None)
    assert calls[0].endswith("mtseisan2026091800.xls")
    assert calls[1].endswith("20260925_mtcurrent.xlsx")
    assert out["fetched"] == [] and len(out["gaps"]) == 2


def test_public_stdout_does_not_contain_authenticated_ledger_or_original_rows(monkeypatch, capsys):
    monkeypatch.setenv("ARGUS_ADMIN_TOKEN", "fixture-secret")
    monkeypatch.setattr(jw, "collect", lambda _: {"fetched": ["2026-09-25"], "gaps": [], "rows": [{"value": 123, "seriesId": "credit.short_balance", "periodEnd": "2026-09-25"}], "csv": "private-csv"})
    monkeypatch.setattr(jw, "import_rows", lambda *a, **k: {"ok": True, "ledger": {"private": "protected-payload"}})
    assert jw.main(["--import"]) == 0
    output = capsys.readouterr().out
    assert "protected-payload" not in output and "fixture-secret" not in output and "private-csv" not in output
    assert output.strip() == '{"ok": true, "stage": "complete"}'


def test_workflow_retries_publication_day_and_does_not_upload_protected_readback():
    from pathlib import Path
    text = Path(".github/workflows/jpx-credit-weekly.yml").read_text()
    assert "cron: '45 7-14 * * 1-5'" in text
    assert "cancel-in-progress: false" in text
    assert "actions/upload-artifact" not in text and "$RUNNER_TEMP/jpx-credit-summary.json" in text


def credit_csv():
    return jw.rows_to_csv(jw.build_rows(jw.parse_sheet(GRID), url="https://example.test/workbook.xls", sha256="ab"*32, observed_at="2026-09-08T00:00:00Z"))


def confirmed_readback():
    return {r["seriesId"]: {"periodEnd": r["periodEnd"], "latestValue": r["value"], "availableFrom": r["availableFrom"]}
            for r in jw.build_rows(jw.parse_sheet(GRID), url="https://example.test/workbook.xls", sha256="ab"*32, observed_at="2026-09-08T00:00:00Z")}


def test_import_requires_values_and_actual_availability_to_match(monkeypatch):
    monkeypatch.setattr(jw, "post_json", lambda *a, **k: {"ok": True})
    for field, value in (("latestValue", 1), ("availableFrom", "2026-09-07T00:00:00Z")):
        wrong = confirmed_readback()
        wrong["credit.long_balance"][field] = value
        monkeypatch.setattr(jw, "ledger_newest_credit", lambda backend, *, token: wrong)
        assert jw.import_rows(credit_csv(), backend="https://backend.test", token="t", expected_newest="2026-08-28")["stage"] == "readback_content"


def test_week_end_uses_thursday_if_friday_is_closed_and_range_is_bounded():
    import pytest
    assert jw.fridays_after("2026-03-13", today=date(2026, 3, 20)) == ["2026-03-19"]
    assert jw.fridays_after("2026-03-19", today=date(2026, 3, 20)) == []
    with pytest.raises(ValueError, match="jpx_period_range_invalid"):
        jw.fridays_after("2000-01-01", today=date(2026, 9, 25))


def test_script_resolves_shared_calendar_when_launched_outside_repo(tmp_path):
    import subprocess
    import sys
    from pathlib import Path
    script = Path(jw.__file__).resolve()
    code = "import runpy; d=runpy.run_path(" + repr(str(script)) + "); from datetime import date; assert d['fridays_after']('2026-03-13',today=date(2026,3,20))==['2026-03-19']"
    subprocess.run([sys.executable, "-c", code], cwd=tmp_path, check=True, capture_output=True, text=True)


def test_listing_finds_renamed_current_workbook_and_excludes_pdf_general_or_offsite():
    import pytest
    html = '''<a href="tvdivq0000001rk9-att/new-credit.xlsx"><span>2026年9月25日申込現在</span></a>
    <a href="tvdivq0000001rk9-att/new-credit.pdf">2026年9月25日申込現在</a>
    <a href="tvdivq0000001rk9-att/mtgaisan2026091800.xls">2026年9月18日申込現在</a>
    <a href="https://other.test/20260925_mtcurrent.xlsx">2026年9月25日申込現在</a>'''
    assert jw.workbook_links(html) == {'2026-09-25': jw.LISTING_URL.rsplit('/', 1)[0] + '/tvdivq0000001rk9-att/new-credit.xlsx'}
    with pytest.raises(ValueError, match='workbooks_missing'):
        jw.workbook_links('<a href="https://other.test/data.xlsx">2026年9月25日</a>')
    with pytest.raises(ValueError, match='period_ambiguous'):
        jw.workbook_links(html + '<a href="tvdivq0000001rk9-att/20260925_mtcurrent.xlsx">2026年9月25日</a>')


def test_publication_deadline_uses_actual_tse_sessions_and_does_not_claim_unknown_calendar():
    from datetime import datetime
    assert jw.publication_due('2026-10-02', datetime.fromisoformat('2026-10-06T15:59:59+09:00')) is False
    assert jw.publication_due('2026-10-02', datetime.fromisoformat('2026-10-06T16:00:00+09:00')) is True
    # Monday is Sports Day; Tuesday and Wednesday are the first two sessions.
    assert jw.publication_due('2026-10-09', datetime.fromisoformat('2026-10-13T17:00:00+09:00')) is False
    assert jw.publication_due('2026-10-09', datetime.fromisoformat('2026-10-14T16:00:00+09:00')) is True
    assert jw.publication_due('2015-10-02', datetime.fromisoformat('2026-10-06T16:00:00+09:00')) is None


def test_listed_missing_or_unreadable_workbook_is_failure_but_not_yet_published_is_pending(monkeypatch):
    listed = lambda: {'2026-09-25': 'https://www.jpx.co.jp/renamed.xlsx'}
    missing = lambda _: None
    result = jw.collect('2026-09-18', today=date(2026, 9, 25), fetcher=missing,
                        listing_fetcher=listed, now_iso='2026-09-25T00:00:00Z')
    assert result['failures'] == ['2026-09-25:published_workbook_unavailable']
    pending = jw.collect('2026-09-25', today=date(2026, 10, 2), fetcher=missing,
                        listing_fetcher=lambda: {}, now_iso='2026-10-05T14:00:00Z')
    assert pending['failures'] == [] and pending['gaps'] == ['2026-10-02']
    overdue = jw.collect('2026-09-25', today=date(2026, 10, 2), fetcher=missing,
                        listing_fetcher=lambda: {}, now_iso='2026-10-06T07:00:00Z')
    assert overdue['failures'] == ['2026-10-02:published_workbook_unavailable']
    monkeypatch.setattr(jw, 'load_workbook_grid', lambda _: [['unexpected shape']])
    changed = jw.collect('2026-09-18', today=date(2026, 9, 25), fetcher=lambda _: b'data',
                         listing_fetcher=listed, now_iso='2026-09-25T00:00:00Z')
    assert changed['failures'] == ['2026-09-25:workbook_unreadable']


def test_collector_follows_listing_link_and_preserves_fetch_availability(monkeypatch):
    requests = []
    renamed = jw.LISTING_URL.rsplit('/', 1)[0] + '/tvdivq0000001rk9-att/revised-name.xlsx'
    monkeypatch.setattr(jw, 'load_workbook_grid', lambda _: current_grid())
    out = jw.collect('2026-09-18', today=date(2026, 9, 25),
                     fetcher=lambda url: requests.append(url) or b'workbook',
                     listing_fetcher=lambda: {'2026-09-25': renamed}, now_iso='2026-09-29T07:01:00Z')
    assert requests == [renamed] and out['failures'] == []
    assert all(r['availableFrom'] == '2026-09-29T07:01:00Z' and r['publishedAt'] == '' for r in out['rows'])


def test_published_gaps_and_listing_failure_cannot_exit_success(monkeypatch, capsys):
    monkeypatch.setattr(jw, 'collect', lambda _: {'fetched': [], 'gaps': ['2026-09-25'], 'failures': ['published'], 'rows': [], 'csv': ''})
    assert jw.main(['--import']) == 1
    assert 'published_weeks_incomplete' in capsys.readouterr().out
    def fail(_):
        raise ValueError('private response must not be logged')
    monkeypatch.setattr(jw, 'collect', fail)
    assert jw.main(['--import']) == 1
    output = capsys.readouterr().out
    assert '"stage": "acquisition"' in output and 'private response' not in output


def test_listing_byte_limit_is_checked_before_parse(monkeypatch):
    import pytest
    monkeypatch.setattr(jw.urllib.request, 'urlopen', lambda *a, **k: io.BytesIO(b'x' * (jw.MAX_LISTING_BYTES + 1)))
    with pytest.raises(ValueError, match='listing_too_large'):
        jw.fetch_listing()
